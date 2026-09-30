"""Build the submission deck from the official template:
    submission/<College>_<Team>_Submission.pptx  (+ a PDF export if LibreOffice is available)

Team details come from submission/team.json, numbers from results/*/summary.json and figures from docs/img.
"""
from __future__ import annotations

import json
import subprocess
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
TEAM = json.loads((ROOT / "submission" / "team.json").read_text())
TEMPLATE = ROOT / "submission" / "templates" / "CollegeName_TeamName_Submission_template.pptx"
OUT = ROOT / "submission" / f"{TEAM['college']}_{TEAM['team_name']}_Submission.pptx"
IMG = ROOT / "docs" / "img"

PURPLE, TITLE, INK, MUTED = RGBColor(0x6D, 0x28, 0xD9), RGBColor(0x70, 0x4E, 0xA6), RGBColor(0x14, 0x14, 0x2B), \
    RGBColor(0x63, 0x63, 0x7E)
LIGHT, LINE, WHITE = RGBColor(0xF4, 0xF1, 0xFD), RGBColor(0xD9, 0xD4, 0xF0), RGBColor(0xFF, 0xFF, 0xFF)
BLUE, GREEN, AMBER, RED = RGBColor(0x2F, 0x80, 0xED), RGBColor(0x1E, 0x9E, 0x6A), RGBColor(0xD9, 0x8E, 0x04), \
    RGBColor(0xC0, 0x39, 0x2B)
FONT = "Calibri"


def load(split: str) -> dict:
    p = ROOT / "results" / split / "summary.json"
    return json.loads(p.read_text()) if p.exists() else {}


RES = {s: load(s) for s in ("test", "dev", "devb", "devc")}
ARCHIVE = ROOT / "results" / "archive" / "devc_heldout_run_engine_v1_duplexrag" / "metrics.json"
V1 = json.loads(ARCHIVE.read_text()) if ARCHIVE.exists() else {}


def m(split: str, cfg: str, key: str, fmt: str = "{:.0f}", suffix: str = "") -> str:
    v = RES.get(split, {}).get(cfg, {}).get(key)
    return "n/a" if v is None else fmt.format(v) + suffix


# ------------------------------------------------------------------ drawing helpers
def text(slide, x, y, w, h, paras, anchor=MSO_ANCHOR.TOP, margin=0.05):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, Inches(margin))
    first = True
    for p in paras:
        if isinstance(p, str):
            p = {"t": p}
        para = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        para.alignment = p.get("align", PP_ALIGN.LEFT)
        para.space_after = Pt(p.get("after", 4))
        runs = p["t"] if isinstance(p["t"], list) else [(p["t"], {})]
        prefix = p.get("bullet")
        if prefix:
            runs = [(prefix + " ", {"color": p.get("bullet_color", PURPLE), "bold": True})] + runs
        for rt, ro in runs:
            r = para.add_run()
            r.text = rt
            f = r.font
            f.name = FONT
            f.size = Pt(ro.get("size", p.get("size", 16)))
            f.bold = ro.get("bold", p.get("bold", False))
            f.italic = ro.get("italic", False)
            f.color.rgb = ro.get("color", p.get("color", INK))
    return tb


def rect(slide, x, y, w, h, fill=LIGHT, line=LINE, radius=True, lw=1.0):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                                 Inches(x), Inches(y), Inches(w), Inches(h))
    if radius:
        shp.adjustments[0] = 0.08
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(lw)
    shp.shadow.inherit = False
    shp.text_frame.text = ""
    return shp


def card(slide, x, y, w, h, title, lines, accent=PURPLE, size=13, title_size=15, fill=LIGHT):
    rect(slide, x, y, w, h, fill=fill, line=LINE)
    bar = rect(slide, x, y + 0.12, 0.07, h - 0.24, fill=accent, line=None, radius=False)
    paras = [{"t": title, "size": title_size, "bold": True, "color": accent, "after": 6}]
    paras += [{"t": ln, "size": size, "bullet": "•", "bullet_color": accent, "after": 3} if isinstance(ln, str)
              else ln for ln in lines]
    text(slide, x + 0.2, y + 0.1, w - 0.3, h - 0.2, paras)
    return bar


def kpi(slide, x, y, w, h, value, label, color=PURPLE):
    rect(slide, x, y, w, h, fill=WHITE, line=LINE)
    text(slide, x + 0.1, y + 0.08, w - 0.2, h * 0.55, [{"t": value, "size": 26, "bold": True, "color": color}])
    text(slide, x + 0.1, y + h * 0.55, w - 0.2, h * 0.45, [{"t": label, "size": 11.5, "color": MUTED}])


def table(slide, x, y, w, rows, col_w, size=12, row_h=0.42, header_fill=PURPLE):
    shp = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w),
                                 Inches(row_h * len(rows)))
    tbl = shp.table
    for j, cw in enumerate(col_w):
        tbl.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        tbl.rows[i].height = Inches(row_h)
        for j, val in enumerate(row):
            cell = tbl.cell(i, j)
            cell.margin_left = cell.margin_right = Inches(0.08)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = True
            para = tf.paragraphs[0]
            color = INK
            bold = i == 0
            if isinstance(val, tuple):
                val, color = val[0], val[1]
                bold = True
            r = para.add_run()
            r.text = str(val)
            r.font.name = FONT
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.color.rgb = WHITE if i == 0 else color
            cell.fill.solid()
            cell.fill.fore_color.rgb = header_fill if i == 0 else (LIGHT if i % 2 == 0 else WHITE)
    return tbl


def clear_body(slide):
    for shp in list(slide.shapes):
        if shp.is_placeholder and shp.placeholder_format.type != 1 and not shp.text_frame.text.strip():
            shp._element.getparent().remove(shp._element)


def picture(slide, path: Path, x, y, w=None, h=None):
    if not path.exists():
        rect(slide, x, y, w or 6, h or 3, fill=LIGHT)
        text(slide, x, y, w or 6, h or 3, [{"t": f"[{path.name} missing]", "color": MUTED}], anchor=MSO_ANCHOR.MIDDLE)
        return
    kw = {}
    if w:
        kw["width"] = Inches(w)
    if h:
        kw["height"] = Inches(h)
    slide.shapes.add_picture(str(path), Inches(x), Inches(y), **kw)


# ------------------------------------------------------------------ slides
def slide_title(s):
    members = [mm for mm in TEAM.get("members", []) if mm.get("name")]
    lines = {
        "Theme ID -": f"Theme ID - {TEAM['theme_id']} ({TEAM['theme_title']})",
        "Team Name -": f"Team Name - {TEAM['team_name']}",
        "College Name -": f"College Name - {TEAM['college']}",
        "Submission Github link -": f"Submission Github link - {TEAM.get('github_url') or '(add repo link)'}",
    }
    for k in range(4):
        mm = TEAM["members"][k] if k < len(TEAM.get("members", [])) else {}
        val = ", ".join(x for x in (mm.get("name", ""), mm.get("email", "")) if x)
        lines[f"Member Name & Email {k + 1}-"] = f"Member Name & Email {k + 1}- {val}"
    for shp in s.shapes:
        if shp.has_text_frame and "Theme ID" in shp.text_frame.text:
            shp.width = Inches(6.85)
            for p in shp.text_frame.paragraphs:
                key = p.text.strip()
                if key in lines and p.runs:
                    p.runs[0].text = lines[key]
                    for r in p.runs[1:]:
                        r.text = ""
                for r in p.runs:
                    r.font.size = Pt(13)
    text(s, 0.85, 6.9, 6.6, 0.4, [{"t": [("DuplexRAG", {"bold": True, "color": PURPLE}),
                                          (" - retrieve while the user is still speaking", {"color": MUTED})],
                                    "size": 13}])
    _ = members


def slide_theme(s):
    clear_body(s)
    text(s, 0.92, 1.55, 6.3, 0.5, [{"t": "Theme 04 · Streaming Live RAG", "size": 20, "bold": True, "color": PURPLE}])
    text(s, 0.92, 2.1, 6.2, 4.6, [
        {"t": "In a full-duplex voice conversation people speak one natural request, not a tidy query.", "size": 16,
         "after": 10},
        {"t": [("Hidden intents - ", {"bold": True}), ("one utterance often carries several questions.", {})],
         "bullet": "•", "size": 15},
        {"t": [("Latency - ", {"bold": True}), ("waiting for the end of speech before searching leaves dead air.", {})],
         "bullet": "•", "size": 15},
        {"t": [("Late details - ", {"bold": True}), ("\"actually it's fifty people\" must refine the answer, not restart it.", {})],
         "bullet": "•", "size": 15},
        {"t": [("Not every turn needs search - ", {"bold": True}), ("\"repeat that in two bullets\" is formatting.", {})],
         "bullet": "•", "size": 15},
        {"t": [("Trust - ", {"bold": True}), ("every claim must be grounded in the provided corpus, and gaps must be said out loud.", {})],
         "bullet": "•", "size": 15},
        {"t": "Goal: a cheap, CPU-only engine that retrieves during speech, answers every hidden intent with "
              "citations, and keeps the answer alive across the conversation.", "size": 15, "bold": True,
         "color": PURPLE, "after": 0},
    ])
    rect(s, 7.45, 1.6, 5.0, 5.3, fill=LIGHT, line=LINE)
    text(s, 7.65, 1.7, 4.6, 0.4, [{"t": "One utterance, three hidden questions", "size": 14, "bold": True,
                                    "color": MUTED}])
    chunks = [("0.0 s", "\"I need to plan a customer workshop in…\""), ("0.8 s", "\"…Pune for 30 people, and I need…\""),
              ("1.6 s", "\"…the cancellation policy and the catering options.\"")]
    for k, (t, c) in enumerate(chunks):
        rect(s, 7.65, 2.15 + k * 0.62, 4.6, 0.52, fill=WHITE, line=LINE)
        text(s, 7.72, 2.17 + k * 0.62, 4.5, 0.5, [{"t": [(t + "  ", {"color": MUTED, "size": 11}),
                                                          (c, {"size": 13})]}], anchor=MSO_ANCHOR.MIDDLE)
    for k, (lab, col) in enumerate([("venue capacity for 30", BLUE), ("cancellation terms", GREEN),
                                    ("catering options", AMBER)]):
        r = rect(s, 7.65 + k * 1.55, 4.1, 1.45, 0.5, fill=col, line=None)
        text(s, 7.65 + k * 1.55, 4.1, 1.45, 0.5, [{"t": lab, "size": 11.5, "bold": True, "color": WHITE,
                                                   "align": PP_ALIGN.CENTER}], anchor=MSO_ANCHOR.MIDDLE)
    text(s, 7.65, 4.8, 4.6, 2.0, [
        {"t": [("then: ", {"color": MUTED}), ("\"Actually it's going to be fifty.\"", {"italic": True})], "size": 13},
        {"t": "→ refine the capacity claim, keep the rest (v1 → v2)", "size": 13, "color": PURPLE, "bold": True,
         "after": 8},
        {"t": [("then: ", {"color": MUTED}), ("\"Can you repeat that in two bullets?\"", {"italic": True})], "size": 13},
        {"t": "→ no retrieval, same citations", "size": 13, "color": PURPLE, "bold": True},
    ])


def slide_gaps(s):
    clear_body(s)
    rows = [["Approach", "How it works", "Gap for full-duplex voice"],
            ["Turn-based RAG", "waits for end of speech, one query for the whole utterance",
             "dead air while searching; compound requests get one blurred query"],
            ["LLM query rewriting / multi-query (RAG-Fusion, HyDE)", "an LLM rewrites or splits the query after the user stops",
             "adds an LLM call to every turn: more latency and cost, and it can invent queries"],
            ["Agentic RAG (plan → search → read loops)", "multi-step tool use by an LLM agent",
             "heavy orchestration and unpredictable latency - the brief asks for simple and cheap"],
            ["Chat-history RAG", "concatenates history into the next query",
             "late details restart the search; context drifts; old citations are lost"],
            ["Intent-based voice assistants", "classify intent → fixed skill",
             "no grounding in documents, no citations, no answer for long-tail policy questions"]]
    table(s, 0.92, 1.7, 11.5, rows, [3.3, 3.9, 4.3], size=12.5, row_h=0.72)
    rect(s, 0.92, 6.25, 11.5, 0.72, fill=PURPLE, line=None)
    text(s, 1.1, 6.25, 11.2, 0.72, [{"t": "The gap: nothing retrieves during speech, refines the answer in place and "
                                          "proves every claim - on a CPU budget, without an LLM in the loop.",
                                     "size": 15, "bold": True, "color": WHITE}], anchor=MSO_ANCHOR.MIDDLE)


def slide_solution(s):
    clear_body(s)
    picture(s, IMG / "architecture.png", 0.92, 1.45, w=11.5)
    text(s, 0.92, 6.9, 11.5, 0.45, [{"t": "Event-driven: on_chunk() per ASR partial, end_turn() at end of speech; "
                                          "retrieval jobs run in the background while the user keeps talking.",
                                     "size": 13, "color": MUTED}])


def slide_demo(s):
    clear_body(s)
    picture(s, IMG / "ui.png", 0.92, 1.6, w=7.6)
    steps = [
        ("Stream", "replay a scenario, type, or speak (Chrome mic) - chunks arrive every ~1 s"),
        ("Watch the controller", "WAIT / RETRIEVE / SUPPRESS per chunk; provisional and multi-intent searches fire mid-sentence"),
        ("Grounded answer", "one line per hidden intent, every claim cited; click a citation to see the source section"),
        ("Late detail", "\"actually fifty people\" → v2 with a diff: added / retained / retired claims, delta search only"),
        ("Format request", "\"two bullets\" → no retrieval, same citations; the trace and cost are downloadable"),
    ]
    paras = []
    for k, (a, b) in enumerate(steps, 1):
        paras.append({"t": [(f"{k}. {a}  ", {"bold": True, "color": PURPLE}), (b, {})], "size": 13, "after": 8})
    video = TEAM.get("video_url") or "submission/DuplexRAG_demo.mp4"
    paras.append({"t": [("Demo video: ", {"bold": True}), (video, {"color": BLUE})], "size": 13})
    text(s, 8.75, 1.6, 3.75, 5.4, paras)


def slide_stack(s):
    clear_body(s)
    cards = [
        ("Models (CPU, ONNX)", ["BAAI bge-small-en-v1.5 - 33M bi-encoder", "ms-marco-MiniLM-L-6-v2 - 22M cross-encoder",
                                "fastembed + ONNX Runtime, 4 threads", "no GPU, no API keys, runs offline"], PURPLE),
        ("Retrieval", ["Okapi BM25 (NumPy, in-house)", "dense cosine search", "reciprocal-rank fusion (k=60)",
                       "cross-encoder rerank + sentence scoring"], BLUE),
        ("Language layer", ["rule-based spoken-language decomposer", "turn gate: rules + NumPy softmax classifier",
                            "self-repair, spoken numbers, anaphora", "optional OpenAI-compatible LLM (verified)"], GREEN),
        ("Serving & UI", ["FastAPI + WebSocket (uvicorn)", "vanilla JS live UI, Web Speech API mic",
                          "CLI: index / serve / replay / ask / bench", "JSONL telemetry + JSON Schema"], AMBER),
        ("Engineering", ["Python 3.12, uv lockfile", "Docker Compose (one command)", "pytest + smoke test (G1)",
                         "git, periodic commits"], PURPLE),
        ("Evaluation", ["discrete-event streaming replay", "4 independently written benchmark splits",
                        "baseline + 7 ablations, gates G1-G6", "matplotlib figures, Playwright demo capture"], BLUE),
    ]
    for k, (t, lines, col) in enumerate(cards):
        x = 0.92 + (k % 3) * 3.9
        y = 1.6 + (k // 3) * 2.75
        card(s, x, y, 3.7, 2.55, t, lines, accent=col, size=12.5)


def slide_impact(s):
    clear_body(s)
    use = [
        ("Voice assistants", "answer 'how do I…' questions from manuals, policies and settings docs while the user is "
                             "still talking - with a source for every claim", PURPLE),
        ("Contact-centre agent assist", "surface the right policy clauses during a live call; refine as the caller adds "
                                        "details; no dead air", BLUE),
        ("Enterprise helpdesk (HR / IT / travel)", "compound, messy questions answered in one turn with explicit "
                                                    "'not in the policy' flags", GREEN),
        ("Hands-free field work", "technicians and drivers ask by voice; CPU-only, offline-capable, cheap to run", AMBER),
    ]
    for k, (t, d, col) in enumerate(use):
        y = 1.6 + k * 1.33
        card(s, 0.92, y, 6.6, 1.2, t, [{"t": d, "size": 12.5, "after": 0}], accent=col, title_size=14)
    lead = RES.get("test", {}).get("duplexrag", {}).get("retrieval_lead_ms_median")
    t_on = RES.get("test", {}).get("duplexrag", {}).get("ttft_ms_p50")
    t_off = RES.get("test", {}).get("no_speculation", {}).get("ttft_ms_p50")
    kpi(s, 7.85, 1.6, 2.2, 1.45, f"{lead / 1000:.1f} s" if lead else "n/a", "median head start: retrieval begins "
        "before the user stops talking")
    kpi(s, 10.25, 1.6, 2.2, 1.45, f"{t_on:.0f} ms" if t_on else "n/a", "median time to first token after speech")
    kpi(s, 7.85, 3.25, 2.2, 1.45, f"{t_off:.0f} ms" if t_off else "n/a",
        "same stack without speculation", color=MUTED)
    kpi(s, 10.25, 3.25, 2.2, 1.45, m("test", "duplexrag", "cost_usd_per_turn_mean", "${:.5f}"),
        "compute cost per turn (CPU only)")
    kpi(s, 7.85, 4.9, 2.2, 1.45, m("test", "duplexrag", "G4_claim_support_pct", "{:.0f}", "%"),
        "claims supported by their citation")
    kpi(s, 10.25, 4.9, 2.2, 1.45, m("test", "duplexrag", "G4_fabricated_ids", "{:.0f}"), "fabricated citation IDs")
    text(s, 7.85, 6.45, 4.6, 0.6, [{"t": "held-out test split, laptop CPU (Ryzen 5 4600H)", "size": 11, "color": MUTED}])


def shrink_title(s, size=34):
    for shp in s.shapes:
        if shp.is_placeholder and shp.placeholder_format.type == 1:
            for p in shp.text_frame.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(size)


def slide_results(s):
    clear_body(s)
    shrink_title(s)

    def row(label, key, fmt="{:.0f}", suffix="%", target=""):
        return [label, m("test", "duplexrag", key, fmt, suffix), m("test", "baseline", key, fmt, suffix), target]

    rows = [["Held-out test (30 sessions, 74 turns)", "DuplexRAG", "Baseline", "Target"],
            row("G2 retrieval starts before end of speech", "G2_early_retrieval_pct", target=">= 80%"),
            row("G2 false triggers on no-retrieval turns", "G2_false_trigger_pct", target="low"),
            row("G3 compound turns: >= 2 intents isolated", "G3_multi_intent_pct", target=">= 70%"),
            row("G4 claims supported by citation", "G4_claim_support_pct", target=">= 85%"),
            row("G4 fabricated citation IDs", "G4_fabricated_ids", "{:.0f}", "", "0"),
            row("G5 refinements: versioned, delta-only", "G5_refinement_pass_pct", target="pass"),
            row("G6 telemetry trace coverage", "G6_trace_coverage_pct", target="100%"),
            row("Retrieval hit@3 (gold section)", "retrieval_hit_at_3_pct"),
            row("Key facts present in answer", "key_fact_recall_pct"),
            row("TTFT p50 after end of speech", "ttft_ms_p50", "{:.0f}", " ms")]
    table(s, 0.92, 1.55, 7.2, rows, [3.75, 1.2, 1.15, 1.1], size=11.5, row_h=0.43)
    text(s, 0.92, 6.35, 7.2, 0.8, [{"t": "Baseline = turn-based RAG: waits for end of speech, one dense query, no rerank, "
                                        "no session memory. G1: one-command Docker build + automated replay; smoke "
                                        "test passes from a clean checkout.", "size": 10.5, "color": MUTED}])
    card(s, 8.4, 1.55, 4.05, 2.85, "Innovation highlights", [
        "speculative per-clause retrieval with entity-aware reuse",
        "refine-in-place answer versions with claim diffs",
        "grounded by construction; claim-level verifier",
        "entity x aspect gaps said out loud",
        "no LLM needed: about $0.0001 per turn"], size=12)
    card(s, 8.4, 4.55, 4.05, 2.55, "Limitations", [
        "synthetic corpus (official corpus unavailable)",
        "extractive wording is less fluent than an LLM",
        "conservative 'not found': some unanswerables still get a related, cited answer",
        "speech simulated from transcripts (live mic in Chrome)"], accent=AMBER, size=11.5)


def slide_next(s):
    clear_body(s)
    steps = [
        ("Plug in the official corpus + real ASR partials", "streaming Whisper / on-device ASR; tune the stability gate on "
                                                            "real partial-hypothesis revisions"),
        ("Learned stability controller", "train the retrieve / wait policy from ASR confidence, prosody and pauses"),
        ("Small verified on-device LLM", "fluent phrasing through the existing claim verifier; stream into TTS"),
        ("Multilingual and code-mixed speech", "Hindi-English, Korean; multilingual bi-encoder and cross-encoder"),
        ("NPU / quantised models", "int8 cross-encoder and cached sentence scores for sub-100 ms p95 on phones"),
        ("Barge-in aware answers (with Theme 5)", "stop, refine and resume the spoken answer when the user interrupts"),
    ]
    for k, (a, b) in enumerate(steps):
        x = 0.92 + (k % 2) * 5.85
        y = 1.6 + (k // 2) * 1.75
        rect(s, x, y, 5.6, 1.55, fill=LIGHT, line=LINE)
        circ = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + 0.2), Inches(y + 0.4), Inches(0.7), Inches(0.7))
        circ.fill.solid()
        circ.fill.fore_color.rgb = PURPLE
        circ.line.fill.background()
        text(s, x + 0.2, y + 0.4, 0.7, 0.7, [{"t": str(k + 1), "size": 18, "bold": True, "color": WHITE,
                                              "align": PP_ALIGN.CENTER}], anchor=MSO_ANCHOR.MIDDLE)
        text(s, x + 1.1, y + 0.15, 4.35, 1.3, [{"t": a, "size": 15, "bold": True, "color": PURPLE, "after": 4},
                                              {"t": b, "size": 12.5}])


def slide_brownie(s):
    clear_body(s)
    lead = RES.get("test", {}).get("duplexrag", {}).get("retrieval_lead_ms_median")
    tiles = [
        ("Retrieves while you talk", f"median {lead / 1000:.1f} s head start; searches fire mid-sentence and are reused"
                                     if lead else "searches fire mid-sentence and are reused", BLUE),
        ("Zero fabricated citations", f"{m('test', 'duplexrag', 'G4_claim_support_pct', '{:.0f}', '%')} of claims "
                                      "verified against the cited section; IDs cannot be invented", GREEN),
        ("CPU-only, about $0.0001 per turn", "two small ONNX models, no GPU, no API keys, works offline in Docker", PURPLE),
        ("Refine, don't restart", "versioned answers with added / retained / retired claims; only the delta is searched",
         BLUE),
        ("Says what it doesn't know", "\"catering for Hinjewadi Tech Park could not be verified\" - entity x aspect "
                                      "coverage checks", AMBER),
        ("Measured honestly", "discrete-event replay; each held-out set written independently after a freeze; "
                              "earlier held-out runs are reported too", GREEN),
    ]
    for k, (t, d, col) in enumerate(tiles):
        x = 0.92 + (k % 3) * 3.9
        y = 1.6 + (k // 3) * 2.75
        card(s, x, y, 3.7, 2.55, t, [{"t": d, "size": 13.5, "after": 0}], accent=col, title_size=16)


def slide_checklist(s):
    body = next(shp for shp in s.shapes if shp.is_placeholder and shp.placeholder_format.type != 1)
    repo = TEAM.get("github_url") or "(add public repo link)"
    video = TEAM.get("video_url") or "submission/DuplexRAG_demo.mp4 in the repo"
    answers = [f"Working prototype code — public or shared GitHub repo (Y): {repo}",
               "README with reproducible setup instructions (Y): README.md, Docker Compose, smoke test",
               f"Demo video, max 5 minutes (Y): {video}",
               f"Presentation file (PPT or PDF) (Y): submission/{OUT.name}"]
    for p, a in zip(body.text_frame.paragraphs, answers):
        if p.runs:
            p.runs[0].text = a
            for r in p.runs[1:]:
                r.text = ""
            for r in p.runs:
                r.font.size = Pt(20)


def main() -> None:
    prs = Presentation(TEMPLATE)
    sl = prs.slides
    slide_title(sl[0])
    slide_theme(sl[1])
    slide_gaps(sl[2])
    slide_solution(sl[3])
    slide_demo(sl[4])
    slide_stack(sl[5])
    slide_impact(sl[6])
    slide_results(sl[7])
    slide_next(sl[8])
    slide_brownie(sl[9])
    slide_checklist(sl[10])
    prs.save(OUT)
    print("wrote", OUT)
    try:
        subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(OUT.parent), str(OUT)],
                       check=True, capture_output=True, timeout=180)
        print("wrote", OUT.with_suffix(".pdf"))
    except Exception as exc:  # LibreOffice optional
        print("PDF export skipped:", exc)


if __name__ == "__main__":
    main()
